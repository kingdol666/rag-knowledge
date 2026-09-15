"""Turn an uploaded file into indexable plain text.

Each loader is optional at runtime: if a parser's dependency is missing the
caller gets a clear, actionable message instead of an opaque ImportError, and
every other format keeps working.

Supported: .txt .md .markdown .rst .log .csv .tsv .json .jsonl .html .htm
           .pdf .docx
"""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass

#: Encodings tried in order for text-ish files. CJK corpora are commonly
#: GB18030/GBK on Windows, so a UTF-8-only reader would garbage them.
_TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "big5", "latin-1")

_TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".log", ".text", ".org"}
_TABLE_EXT = {".csv", ".tsv"}
_JSON_EXT = {".json", ".jsonl", ".ndjson"}
_HTML_EXT = {".html", ".htm", ".xhtml"}

SUPPORTED_EXTENSIONS = sorted(
    _TEXT_EXT | _TABLE_EXT | _JSON_EXT | _HTML_EXT | {".pdf", ".docx"}
)


class UnsupportedFileType(ValueError):
    """Raised when no loader can handle the uploaded file."""


@dataclass(frozen=True)
class Extracted:
    """Result of parsing one uploaded file."""

    text: str
    kind: str
    extension: str
    #: Non-fatal notes worth surfacing to the user (e.g. "3 pages had no text").
    warnings: tuple[str, ...] = ()


def _decode(data: bytes) -> str:
    """Decode bytes using the first encoding that does not fail."""
    for encoding in _TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _load_pdf(data: bytes) -> Extracted:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on install
        raise UnsupportedFileType(
            "PDF support needs the 'pypdf' package. Install it with: "
            "uv pip install --python <venv> pypdf"
        ) from exc

    reader = PdfReader(io.BytesIO(data))
    pages, empty = [], 0
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:  # a single malformed page must not kill the upload
            page_text = ""
        if page_text.strip():
            pages.append(page_text.strip())
        else:
            empty += 1
    warnings = ()
    if empty:
        warnings = (f"{empty}/{len(reader.pages)} page(s) contained no extractable text "
                    "(scanned pages need OCR)",)
    return Extracted(text=_clean("\n\n".join(pages)), kind="pdf",
                     extension=".pdf", warnings=warnings)


def _load_docx(data: bytes) -> Extracted:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - depends on install
        raise UnsupportedFileType(
            "DOCX support needs the 'python-docx' package. Install it with: "
            "uv pip install --python <venv> python-docx"
        ) from exc

    document = docx.Document(io.BytesIO(data))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return Extracted(text=_clean("\n\n".join(parts)), kind="docx", extension=".docx")


def _load_table(text: str, extension: str) -> Extracted:
    delimiter = "\t" if extension == ".tsv" else ","
    rows = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        cells = [c.strip() for c in row if c and c.strip()]
        if cells:
            rows.append(" | ".join(cells))
    return Extracted(text=_clean("\n".join(rows)), kind="table", extension=extension)


def _load_json(text: str, extension: str) -> Extracted:
    """JSON/JSONL -> readable text.

    Objects are flattened to ``key: value`` lines so BM25 and the embedding
    model see the field names, not just the values.
    """
    records: list = []
    if extension in {".jsonl", ".ndjson"}:
        for line in text.splitlines():
            line = line.strip()
            if line:
                try:
                    records.append(json.loads(line))
                except json.JSONDecodeError:
                    records.append(line)
    else:
        try:
            parsed = json.loads(text)
            records = parsed if isinstance(parsed, list) else [parsed]
        except json.JSONDecodeError:
            return Extracted(text=_clean(text), kind="text", extension=extension)

    lines: list[str] = []

    def walk(node, prefix: str = "") -> None:
        if isinstance(node, dict):
            for key, value in node.items():
                walk(value, f"{prefix}{key}: ")
        elif isinstance(node, list):
            for item in node:
                walk(item, prefix)
        else:
            lines.append(f"{prefix}{node}")

    for record in records:
        walk(record)
        lines.append("")
    return Extracted(text=_clean("\n".join(lines)), kind="json", extension=extension)


def _load_html(text: str) -> Extracted:
    text = re.sub(r"(?is)<(script|style|noscript)[^>]*>.*?</\1>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</h[1-6]>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    for entity, char in (("&nbsp;", " "), ("&amp;", "&"), ("&lt;", "<"),
                         ("&gt;", ">"), ("&quot;", '"'), ("&#39;", "'")):
        text = text.replace(entity, char)
    return Extracted(text=_clean(text), kind="html", extension=".html")


def extract(filename: str, data: bytes) -> Extracted:
    """Extract plain text from ``data``, dispatching on ``filename``'s suffix."""
    extension = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if extension == ".pdf":
        return _load_pdf(data)
    if extension == ".docx":
        return _load_docx(data)

    text = _decode(data)
    if extension in _TABLE_EXT:
        return _load_table(text, extension)
    if extension in _JSON_EXT:
        return _load_json(text, extension)
    if extension in _HTML_EXT:
        return _load_html(text)
    if extension in _TEXT_EXT or extension == "":
        return Extracted(text=_clean(text), kind="text", extension=extension or ".txt")
    raise UnsupportedFileType(
        f"unsupported file type '{extension}'. Supported: "
        + ", ".join(SUPPORTED_EXTENSIONS)
    )
