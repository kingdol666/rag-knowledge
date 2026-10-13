# -*- coding: utf-8 -*-
"""Render HTML slides -> PNG (playwright) -> assemble 16:9 PPTX (python-pptx)."""
import pathlib
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

BASE = pathlib.Path(r"D:/codes/ClaudeGPT/rag_project/rag-knowledge/review/deck-benchmark-20261013")
SLIDES = ["slide-01.html", "slide-02.html", "slide-03.html"]
PNG_DIR = BASE / "png"
PNG_DIR.mkdir(exist_ok=True)

with sync_playwright() as p:
    browser = p.chromium.launch()
    page = browser.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=2)
    for name in SLIDES:
        page.goto((BASE / "slides" / name).as_uri())
        page.wait_for_timeout(500)
        out = PNG_DIR / (name.replace(".html", ".png"))
        page.screenshot(path=str(out), clip={"x": 0, "y": 0, "width": 1920, "height": 1080})
        print("shot:", out.name)
    browser.close()

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
for name in SLIDES:
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(PNG_DIR / (name.replace(".html", ".png"))), 0, 0,
                         width=prs.slide_width, height=prs.slide_height)
out_pptx = BASE / "QDCVR-基准测试汇报-20261013.pptx"
prs.save(str(out_pptx))
print("pptx:", out_pptx)
